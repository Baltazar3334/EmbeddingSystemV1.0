import os

from docx import Document
from PyPDF2 import PdfReader


# DOCX =====================================================
def load_contract_from_docx(path: str) -> str:

    doc = Document(path)

    full_text = []

    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()

                if text:
                    full_text.append(text)

    return "\n".join(full_text)


# PDF =====================================================
def load_contract_from_pdf(path: str) -> str:

    reader = PdfReader(path)

    full_text = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            full_text.append(text)

    return "\n".join(full_text)


# UNIVERSAL LOADER =====================================================
def load_contract(path: str) -> str:

    ext = os.path.splitext(path)[1].lower()

    if ext == ".docx":
        return load_contract_from_docx(path)

    elif ext == ".pdf":
        return load_contract_from_pdf(path)

    else:
        raise ValueError(f"Unsupported file type: {ext}")